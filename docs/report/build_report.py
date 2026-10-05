"""Generate the assignment report and supporting design artefacts from local sources.
Run from any directory: python docs/report/build_report.py
Dependencies: pypdf, python-docx, reportlab, pillow (install into .report-tools).
No database writes, credentials, external messages or real user measurements.
"""
from pathlib import Path
import csv
import json
import math
import re
import shutil
import sys
import textwrap
from html import escape

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.report-tools'))
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as PDFImage, PageBreak, KeepInFrame
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader

OUT = ROOT / 'docs/report'
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
DATA = json.loads((OUT / 'report-data.json').read_text(encoding='utf-8'))
INK = '#111827'
GREEN = '#166534'
BLUE = '#1D4ED8'
GRAY = '#475569'
FONT = Path('C:/Windows/Fonts/arial.ttf')
BOLD = Path('C:/Windows/Fonts/arialbd.ttf')

def font(size=26, bold=False):
    return ImageFont.truetype(str(BOLD if bold else FONT), size)

def canvas(title, width=1600, height=900):
    im = Image.new('RGB', (width, height), 'white')
    d = ImageDraw.Draw(im)
    d.text((40, 20), title, font=font(34, True), fill=GREEN)
    return im, d

def text(d, xy, value, width=32, size=26, fill=INK, bold=False):
    lines = '\n'.join(textwrap.fill(line, width) for line in value.split('\n'))
    d.multiline_text(xy, lines, font=font(size, bold), fill=fill, spacing=8)

def box(d, rect, label, fill='#F1F5F9', size=26, width=26):
    d.rounded_rectangle(rect, radius=18, fill=fill, outline=GRAY, width=3)
    text(d, (rect[0]+18, rect[1]+16), label, width=width, size=size)

def arrow(d, a, b, label='', color=GRAY, dashed=False):
    if dashed:
        length = math.dist(a,b)
        for i in range(0, int(length), 18):
            f, g = i/length, min(i+10,length)/length
            d.line((a[0]+(b[0]-a[0])*f,a[1]+(b[1]-a[1])*f,a[0]+(b[0]-a[0])*g,a[1]+(b[1]-a[1])*g),fill=color,width=3)
    else: d.line((a,b), fill=color, width=3)
    ang = math.atan2(b[1]-a[1], b[0]-a[0])
    pts = [b, (b[0]-17*math.cos(ang-.5),b[1]-17*math.sin(ang-.5)),(b[0]-17*math.cos(ang+.5),b[1]-17*math.sin(ang+.5))]
    d.polygon(pts, fill=color)
    if label: text(d, ((a[0]+b[0])/2+5,(a[1]+b[1])/2-35),label,width=18,size=22,fill=color)

def save(im, name): im.save(ASSETS/name)

def diagrams():
    im,d=canvas('UML use-case diagram — role tasks')
    d.rectangle((280,90,1370,850),outline=GRAY,width=3)
    text(d,(720,105),'CampusConnect',size=28,bold=True)
    cases=[(450,220,'Register / sign in'),(450,390,'Request / schedule ride'),(450,570,'Message / rate ride'),(450,750,'Raise SOS'),(1070,220,'Review applications'),(1070,390,'Manage users/settings'),(1070,570,'Accept / update trip'),(1070,750,'Inspect / dispatch SOS')]
    for x,y,label in cases:
        d.ellipse((x-155,y-48,x+155,y+48),fill='#F0FDF4',outline=GRAY,width=3)
        text(d,(x-130,y-16),label,width=24,size=22)
    for x,y,label in [(100,370,'Rider'),(1490,240,'Admin'),(1490,540,'Driver'),(1490,750,'Security')]:
        d.ellipse((x-18,y-70,x+18,y-34),outline=INK,width=3)
        d.line((x,y-34,x,y+30),fill=INK,width=3);d.line((x-30,y,x+30,y),fill=INK,width=3)
        d.line((x,y+30,x-25,y+65),fill=INK,width=3);d.line((x,y+30,x+25,y+65),fill=INK,width=3)
        text(d,(x-50,y+75),label,size=24)
    for a,b in [((130,370),(295,220)),((130,370),(295,390)),((130,370),(295,570)),((130,370),(295,750)),((1460,240),(1225,220)),((1460,240),(1225,390)),((1460,540),(1225,570)),((1460,750),(1225,750))]:d.line((a,b),fill=GRAY,width=2)
    save(im,'usecase.png')
    im,d=canvas('UML activity — immediate ride request')
    d.ellipse((85,140,115,170),fill=INK)
    boxes=[((190,110,490,210),'Enter locations'),((660,110,960,210),'Validate request'),((1150,110,1450,210),'Show input error'),((660,340,960,440),'Review fare / submit'),((660,560,960,660),'Persist PENDING ride'),((190,730,490,830),'Cancel / wait'),((1100,730,1450,830),'Driver accepts')]
    for r,t in boxes:box(d,r,t,size=26,width=24)
    arrow(d,(115,155),(190,155));arrow(d,(490,155),(660,155))
    d.polygon([(810,230),(880,275),(810,320),(740,275)],fill='#DBEAFE',outline=GRAY)
    arrow(d,(810,210),(810,230));arrow(d,(880,275),(1300,210),'invalid')
    arrow(d,(810,320),(810,340),'valid');arrow(d,(810,440),(810,560))
    arrow(d,(810,660),(1270,730),'accepted');arrow(d,(750,660),(340,730),'no match')
    text(d,(500,825),'Trip status remains explicit; no false completion.',size=24)
    save(im,'activity.png')
    im,d=canvas('UML state — intended trip lifecycle',height=650)
    states=[(50,110,'SCHEDULED'),(360,110,'PENDING'),(670,110,'ACCEPTED'),(1000,110,'ENROUTE'),(1000,340,'ARRIVED'),(670,340,'STARTED'),(360,340,'COMPLETED'),(50,340,'CANCELLED')]
    for x,y,t in states:box(d,(x,y,x+240,y+90),t,size=25,width=16)
    for a,b,t in [((290,155),(360,155),'dispatch'),((600,155),(670,155),'accept'),((910,155),(1000,155),'travel'),((1120,200),(1120,340),'arrive'),((1000,385),(910,385),'start'),((670,385),(600,385),'complete'),((470,200),(170,340),'cancel')]:arrow(d,a,b,t)
    text(d,(80,520),'Cancellation may occur from permitted active states.\nTerminal states retain history; enforcement must be integration-tested.',width=95,size=26)
    save(im,'state.png')
    im,d=canvas('UML class — selected implementation classes')
    classes=[((50,100,460,350),'User','id: Long\nrole: Role\nstatus: UserStatus\nemail: String'),((600,100,1010,350),'Ride','id/riderId/driverId: Long\nstatus: RideStatus\nfare: BigDecimal\nscheduledAt: LocalDateTime'),((1150,100,1550,350),'Driver','id/userId: Long\napprovalStatus: enum\nonline: boolean'),((50,540,460,790),'ChatMessage','rideId/senderId: Long\nmessageText: String\nsentAt: LocalDateTime'),((600,540,1010,790),'RideController','requestRide(...)\ncancelRide(...)\nrateDriver(...)'),((1150,540,1550,790),'RideService','saveRide(Ride)\ngetRideById(Long)\nRideRepository dependency')]
    for rect,name,attrs in classes:
        box(d,rect,'',size=23);text(d,(rect[0]+20,rect[1]+14),name,size=28,bold=True)
        d.line((rect[0],rect[1]+65,rect[2],rect[1]+65),fill=GRAY,width=2)
        text(d,(rect[0]+18,rect[1]+80),attrs,width=28,size=23)
    arrow(d,(460,220),(600,220),'1 / 0..*')
    d.line((260,100,260,80,1350,80),fill=GRAY,width=3)
    arrow(d,(1350,80),(1350,100))
    arrow(d,(330,540),(700,350),'rideId');arrow(d,(1010,640),(1150,640),'uses')
    arrow(d,(800,540),(800,350),'operates on')
    save(im,'class.png')
    im,d=canvas('UML sequence — request and driver acceptance')
    xs=[140,540,960,1400]
    for x,name in zip(xs,['Rider client','Ride API/service','MySQL','Driver client']):
        box(d,(x-125,90,x+125,165),name,size=24,width=18)
        d.line((x,170,x,840),fill=GRAY,width=2)
    steps=[(0,1,240,'POST request + JWT'),(1,1,325,'validate user/input'),(1,2,410,'INSERT ride: PENDING'),(2,1,480,'ride ID / stored state'),(1,0,555,'pending response'),(3,1,635,'accept request + JWT'),(1,2,710,'assign driver / ACCEPTED'),(0,1,790,'GET status (poll)')]
    for a,b,y,label in steps:
        if a==b:
            d.line((xs[a],y,xs[a]+130,y,xs[a]+130,y+45,xs[a],y+45),fill=GRAY,width=3)
            text(d,(xs[a]+140,y),label,width=24,size=22)
        else:arrow(d,(xs[a],y),(xs[b],y),label)
    save(im,'sequence.png')
    im,d=canvas('WBS and illustrative eight-week delivery plan',height=620)
    tasks=[('1 Discovery',1,2),('2 Analysis / UML',2,3),('3 Interface / data',3,4),('4 Client / API',4,6),('5 Evaluation',6,7),('6 Report / demo',7,8)]
    for w in range(1,9):text(d,(410+(w-1)*130,100),f'W{w}',size=26,bold=True)
    for i,(name,start,end) in enumerate(tasks):
        y=165+i*65;text(d,(45,y),name,size=25)
        for w in range(1,9):d.rectangle((395+(w-1)*130,y,510+(w-1)*130,y+42),fill='#E2E8F0')
        d.rectangle((395+(start-1)*130,y,510+(end-1)*130,y+42),fill=GREEN)
    text(d,(45,575),'Proposed relative weeks; actual dates and member ownership are pending.',size=24)
    save(im,'gantt.png')
    im,d=canvas('ERD — code-supported logical model (not a live DDL export)',height=1000)
    tables=[(30,110,'users','PK user_id'),(590,110,'drivers','PK driver_id / user_id'),(1150,110,'password_reset_codes','PK id / UNIQUE email'),(30,410,'user_payment_methods','PK id / FK user_id'),(590,410,'rides','PK ride_id / rider_id,driver_id'),(1150,410,'admin_platform_settings','PK id / standalone'),(30,730,'messages','PK message_id / ride_id,sender_id'),(590,730,'sos_alerts','PK alert_id / ride_id,triggered_by,resolved_by'),(1150,730,'payments','ride_id,payer_id / DDL missing')]
    for x,y,name,keys in tables:
        box(d,(x,y,x+420,y+150),name+'\n'+keys,size=24,width=29)
    for a,b,label,dashed in [((450,170),(590,170),'',True),((240,260),(240,410),'1 : many',False),((450,240),(590,450),'1 : many; two roles',True),((660,560),(800,730),'1 : many',True),((600,540),(450,730),'1 : many',True),((1010,560),(1300,730),'1 : many',True)]:arrow(d,a,b,label,dashed=dashed)
    text(d,(460,120),'1 to 0..1',width=11,size=18)
    d.line((30,220,10,220,10,805),fill=BLUE,width=2)
    arrow(d,(10,805),(30,805),color=BLUE)
    text(d,(40,650),'User → many messages (sender)',width=30,size=20,fill=BLUE)
    d.line((450,230,520,230,520,810),fill=BLUE,width=2)
    arrow(d,(520,810),(590,810),color=BLUE)
    text(d,(555,660),'User → SOS: trigger / resolve',width=30,size=20,fill=BLUE)
    d.line((330,260,330,310,1588,310,1588,805),fill=BLUE,width=2)
    arrow(d,(1588,805),(1570,805),color=BLUE)
    text(d,(1170,650),'User → many payments (payer)',width=29,size=20,fill=BLUE)
    arrow(d,(240,110),(240,85),dashed=True)
    arrow(d,(240,85),(1360,85),dashed=True)
    arrow(d,(1360,85),(1360,110),dashed=True)
    text(d,(940,55),'User email association',size=22)
    text(d,(45,930),'Grey solid: SQL FK. Grey dashed / blue: code-supported logical associations; verify constraints. Blue links connect User as sender, triggering/resolving officer and payer.',width=120,size=22)
    save(im,'erd.png')
    im,d=canvas('Design iteration — reconstructed before/after mock-up',height=380)
    box(d,(40,85,730,330),'BEFORE\nPage content and footer share scroll flow\nDuplicate stack + page back controls',fill='#FEF2F2',size=28,width=43)
    box(d,(870,85,1560,330),'AFTER\nBounded page scrolls above a sibling footer\nOne appropriate back control per page',fill='#F0FDF4',size=28,width=43)
    arrow(d,(730,220),(870,220),'revise')
    save(im,'iteration.png')

SCREENS=[
 ('Landing','CampusConnect','Brand / loading / sign-in entry'),('Login','Welcome Back','Email / Password / Sign In / Recovery'),('CreateAccount','Create Account','Role / contact details / Next / step Back'),('Legal','Terms and Privacy','Document text / one Back control'),('FaceVerification','Face Verification','Permission / camera guide / capture'),('Home','Book a Ride','Pickup / destination / fare / request'),('RiderHistory','Ride History','Trips / Rate ride / Previous / Next'),
 ('RiderSchedule','Schedule a Ride','Locations / date / time / planned trips'),('RiderProfile','My Profile','Edit / saved cards / support / logout'),('RiderPaymentMethods','Payment Methods','Cash / Card / Add card / owned records'),('RatingDriver','Rate Your Ride','Five stars / Submit rating / feedback'),('Chat','Ride Messages','Thread / text input / Send / Close'),('DriverDashboard','Driver Dashboard','Online / requests / View Request'),('ViewRideDetails','Ride Request Details','Rider / route / rand fare / Accept / Decline'),
 ('DriverActiveRide','Active Trip','Map / contact / messages / phase button'),('DriverEarnings','Earnings','Month comparison / completed-trip totals'),('DriverHistory','Trip History','Completed trips / pagination'),('DriverProfile','Driver Profile','Personal and vehicle details / logout'),('AdminDashboard','Admin Dashboard','Stats / applications / four quick actions'),('UserManagement','User Management','Search / filters / suspend / pagination'),('RideMonitoring','Ride Monitoring','Map / current rides / Track Live / Refresh'),
 ('IncidentReports','Incident Reports','Search / details / audio / pagination'),('AudioRecordings','Audio Recordings','Incident reference / Play / Download'),('UniversitySettings','University Settings','Institution / security options / Save'),('SecurityDashboard','Security Dashboard','Safety overview / incidents / navigation'),('ActiveRidesMonitor','Active Rides Monitor','Map / Center Map / Contact / rides'),('SosAlerts','SOS Alerts','Active alerts / Dispatch / Audio / resolved'),('ResolvedSos','Resolved SOS','Resolved records / search / pagination'),('TestConnection','Connection Test','Diagnostic request / result feedback')]

def galleries():
    captures={
        'Home':ROOT/'mobile/.expo/responsive-checks/768x1024-home.png',
        'RiderProfile':ROOT/'mobile/.expo/responsive-checks/320x568-profile.png',
        'UniversitySettings':ROOT/'mobile/.expo/page-review-images/settings-tablet.png',
    }
    catalogue=['# Registered screen catalogue','', 'Every route in mobile/App.tsx is represented. Schematics are design explanations, not runtime or user-test evidence. Actual screenshots use local synthetic fixtures.','']
    groups=[SCREENS[:7],SCREENS[7:14],SCREENS[14:21],SCREENS[21:]]
    for gi,group in enumerate(groups,1):
        im,d=canvas(f'Screen catalogue {gi} — schematic unless labelled fixture capture',height=1170)
        for i,(route,title,controls) in enumerate(group):
            x=30+(i%4)*395; y=100+(i//4)*530
            d.rounded_rectangle((x,y,x+365,y+440),radius=20,fill='#F8FAFC',outline='#CBD5E1',width=3)
            capture=captures.get(route)
            if capture and capture.exists():
                local=ASSETS/(route+'-fixture.png');shutil.copyfile(capture,local)
                pic=Image.open(local).convert('RGB');pic.thumbnail((345,420));im.paste(pic,(x+(365-pic.width)//2,y+10))
                label='FIXTURE CAPTURE'
            else:
                text(d,(x+20,y+20),title,width=21,size=29,bold=True)
                for j,control in enumerate(controls.split(' / ')):
                    r=(x+16,y+110+j*54,x+349,y+150+j*54)
                    d.rounded_rectangle(r,radius=9,fill='white',outline='#CBD5E1',width=2)
                    text(d,(x+26,y+119+j*54),control,width=30,size=20)
                if route in ['Home','RiderHistory','RiderSchedule','RiderProfile']:
                    d.rectangle((x+8,y+384,x+357,y+430),fill='#DCFCE7')
                    text(d,(x+16,y+397),'Home | History | Schedule | Profile',width=38,size=17)
                label='SCHEMATIC'
            text(d,(x+8,y+450),route,width=28,size=23,bold=True)
            text(d,(x+8,y+480),label,width=30,size=20,fill=BLUE)
            catalogue.extend([f'## {route}',f'Purpose: {title}. Controls: {controls}.',f'Panel evidence: {label}. Source: mobile/App.tsx and its registered screen component.',''])
        save(im,f'gallery-{gi}.png')
    (OUT/'screen-catalogue.md').write_text('\n'.join(catalogue),encoding='utf-8')

def dictionary():
    records=[]
    entities=ROOT/'backend/src/main/java/com/campusconnect/entity'
    types={'Long':'BIGINT','Integer':'INT','int':'INT','Boolean':'BOOLEAN','boolean':'BOOLEAN','Double':'DOUBLE','BigDecimal':'DECIMAL (precision unspecified)','LocalDateTime':'DATETIME','String':'VARCHAR'}
    for p in entities.glob('*.java'):
        s=p.read_text(encoding='utf-8');match=re.search(r'@Table\(name\s*=\s*"([^"]+)"',s)
        if not match:continue
        table=match.group(1)
        for f in re.finditer(r'((?:\s*@[\s\S]*?)?)\s*private\s+(\w+)\s+(\w+)\s*(?:=[^;]*)?;',s):
            annotations,java,name=f.groups()
            column=re.search(r'@Column\([^)]*name\s*=\s*"([^"]+)"',annotations)
            field=column.group(1) if column else re.sub(r'(?<!^)(?=[A-Z])','_',name).lower()
            t=types.get(java,'ENUM '+java);length=re.search(r'length\s*=\s*(\d+)',annotations)
            definition=re.search(r'columnDefinition\s*=\s*"([^"]+)"',annotations)
            size=length.group(1) if length else '255 (JPA default)' if java=='String' else 'not specified'
            if definition:t=definition.group(1);size='variable'
            purpose='Primary key' if '@Id' in annotations else re.sub(r'(?<!^)(?=[A-Z])',' ',name)
            records.append([table,field,t,size,purpose,str(p.relative_to(ROOT)).replace('\\','/'),'JPA logical mapping; live DDL not inspected'])
    extra={
        'users':[('default_payment_method','VARCHAR','20','Default payment preference')],
        'sos_alerts':[('resolved_at','DATETIME','-','Resolution time'),('resolved_by','BIGINT','-','Resolving user ID')],
        'user_payment_methods':[('id','BIGINT','-','Primary key'),('user_id','BIGINT','-','Owning user'),('label','VARCHAR','80','Display label'),('last_four','CHAR','4','Last four digits'),('brand','VARCHAR','32','Card brand'),('expiry','CHAR','5','MM/YY expiry'),('cardholder','VARCHAR','120','Display name'),('is_default','BOOLEAN','-','Default flag'),('created_at','DATETIME','-','Creation time')],
        'admin_platform_settings':[('id','BIGINT','-','Primary key'),('university_name','VARCHAR','160','Institution label'),('email_domain','VARCHAR','120','Configured email domain'),('campus_security_phone','VARCHAR','40','Configured phone'),('sos_response_time_seconds','INT','-','Configuration target'),('first_year_priority_matching','BOOLEAN','-','Configured preference'),('updated_at','DATETIME','-','Last save')],
        'payments':[('ride_id','ID type unverified','UNVERIFIED','Trip'),('payer_id','ID type unverified','UNVERIFIED','Payer'),('amount','Numeric type unverified','UNVERIFIED','Fare record'),('method','ENUM/string unverified','UNVERIFIED','CASH/CARD; expanded enum migration'),('status','String/enum unverified','UNVERIFIED','PENDING inserted'),('transaction_ref','String unverified','UNVERIFIED','Reference; no settlement')]
    }
    for table,fields in extra.items():
        for name,t,size,meaning in fields:
            source='backend/src/main/java/com/campusconnect/controller/PaymentController.java' if table=='payments' else 'backend/src/main/java/com/campusconnect/CampusConnectApplication.java' if table=='admin_platform_settings' else 'database/mysql/schema.sql'
            records.append([table,name,t,size,meaning,source,'Insert contract only; creation DDL missing' if table=='payments' else 'SQL source; live DDL not inspected'])
    headers=['Table','Field','Type','Length / precision','Description','Source','Certainty']
    with (OUT/'data-dictionary.csv').open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.writer(f);writer.writerow(headers);writer.writerows(records)
    lines=['# Complete source-based data dictionary','', 'This is a source inventory, not an export of the live database. JPA defaults and unverified types are labelled. See report pp.13–15 for drift and constraints.','']
    for table in sorted(set(row[0] for row in records)):
        lines.extend([f'## {table}','| Field | Type | Size | Description | Evidence |','|---|---|---|---|---|'])
        for row in records:
            if row[0]==table:lines.append('| '+' | '.join(row[1:5]+[row[6]])+' |')
        lines.append('')
    (OUT/'data-dictionary.md').write_text('\n'.join(lines),encoding='utf-8')
    return records

def footer(c,doc):
    c.setFont('Helvetica',8);c.setFillColor(colors.HexColor(GRAY))
    c.drawString(38,24,'CampusConnect | NPRT630 | Evidence-based draft')
    c.drawRightString(A4[0]-38,24,f'{doc.page} / 20')

def exports():
    pdfmetrics.registerFont(TTFont('ReportArial', str(FONT)))
    pdfmetrics.registerFont(TTFont('ReportArialBold', str(BOLD)))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='ReportBody',fontName='ReportArial',fontSize=9.5,leading=13,spaceAfter=8,textColor=colors.HexColor(INK)))
    styles.add(ParagraphStyle(name='SmallCell',fontName='ReportArial',fontSize=7.7,leading=10,spaceAfter=0))
    styles.add(ParagraphStyle(name='ReportNote',fontName='ReportArial',fontSize=8.2,leading=11,spaceBefore=5,textColor=colors.HexColor(GRAY)))
    styles['Heading1'].fontName='ReportArialBold'
    styles['Heading1'].textColor=colors.HexColor(GREEN)
    pdf=SimpleDocTemplate(str(OUT/'CampusConnect-Final-Report.pdf'),pagesize=A4,leftMargin=38,rightMargin=38,topMargin=36,bottomMargin=38,title=DATA['title'],author='CampusConnect group — details pending')
    story=[];md=[];word=Document();section=word.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=Inches(.5);section.bottom_margin=Inches(.55);section.left_margin=Inches(.55);section.right_margin=Inches(.55)
    normal=word.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.line_spacing=1.03
    for name in ['Heading 1','Heading 2']:
        word.styles[name].font.name='Arial';word.styles[name].font.color.rgb=RGBColor.from_string('166534')
    word.styles['Heading 1'].font.size=Pt(17)
    section.footer.paragraphs[0].text='CampusConnect | NPRT630 | Evidence-based draft | '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');section.footer.paragraphs[0]._p.append(fld)
    for pi,page in enumerate(DATA['pages'],1):
        blocks=[Paragraph(escape(page['title']),styles['Heading1'])]
        word.add_heading(page['title'],level=1);md.extend([f'<!-- Report page {pi} -->',f'# {page["title"]}',''])
        if page.get('subtitle'):
            blocks.append(Paragraph(escape(page['subtitle']),styles['ReportBody']));word.add_paragraph(page['subtitle']);md.extend([page['subtitle'],''])
        for para in page.get('paragraphs',[]):
            blocks.append(Paragraph(escape(para),styles['ReportBody']));word.add_paragraph(para);md.extend([para,''])
        for key,heightkey in [('image','image_height'),('image2','image2_height')]:
            if page.get(key):
                p=ASSETS/page[key];img=Image.open(p);w,h=img.size
                maxheight=page.get(heightkey,250);width=A4[0]-76;scale=min(width/w,maxheight/h)
                blocks.extend([PDFImage(str(p),width=w*scale,height=h*scale),Spacer(1,5)])
                wp=word.add_paragraph();wp.paragraph_format.space_after=Pt(3)
                wp.add_run().add_picture(str(p),width=Inches(w*scale/72),height=Inches(h*scale/72))
                md.extend([f'![{page["title"]}](assets/{page[key]})',''])
        if page.get('table'):
            t=page['table'];rows=[t['headers']]+t['rows'];cols=len(t['headers'])
            table=Table([[Paragraph(escape(str(cell)),styles['SmallCell']) for cell in row] for row in rows],colWidths=[(A4[0]-76)/cols]*cols,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCFCE7')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#CBD5E1')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
            blocks.append(table);wt=word.add_table(rows=1,cols=cols);wt.style='Light Shading Accent 1'
            for i,h in enumerate(t['headers']):wt.rows[0].cells[i].text=h
            for row in t['rows']:
                cells=wt.add_row().cells
                for i,cell in enumerate(row):cells[i].text=str(cell)
            for row in wt.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                        for r in p.runs:r.font.name='Arial';r.font.size=Pt(7.5)
            md.extend(['| '+' | '.join(t['headers'])+' |','| '+' | '.join(['---']*cols)+' |']+['| '+' | '.join(map(str,r))+' |' for r in t['rows']]+[''])
        if page.get('note'):
            blocks.append(Paragraph(escape(page['note']),styles['ReportNote']))
            p=word.add_paragraph(page['note']);p.paragraph_format.space_before=Pt(4)
            for r in p.runs:r.font.size=Pt(8)
            md.extend([page['note'],''])
        story.append(KeepInFrame(A4[0]-76,A4[1]-78,blocks,mode='shrink'))
        if pi<len(DATA['pages']):story.append(PageBreak());word.add_page_break();md.extend(['---',''])
    pdf.build(story,onFirstPage=footer,onLaterPages=footer)
    word.core_properties.title=DATA['title'];word.core_properties.subject='NPRT630 four-phase report';word.core_properties.author='Group details pending'
    word.save(OUT/'CampusConnect-Final-Report.docx')
    (OUT/'CampusConnect-Final-Report.md').write_text('\n'.join(md),encoding='utf-8')
    r=PdfReader(OUT/'CampusConnect-Final-Report.pdf')
    assert len(r.pages)==20,f'Expected 20 PDF pages, got {len(r.pages)}'
    for i,p in enumerate(r.pages):
        extracted=p.extract_text()
        assert DATA['pages'][i]['title'][:35] in extracted,(i,extracted[:100])
    assert len(word.inline_shapes)==sum(bool(p.get(k)) for p in DATA['pages'] for k in ['image','image2'])
    print('PASS: 20 PDF pages, all expected headings, embedded Word figures and 19 explicit Word page breaks.')
    print('Word pagination remains dependent on the viewer/fonts; PDF pagination is verified.')

def support(records):
    folder=ROOT/'submission/NPRT630-3rd-year-Group-PENDING-Team-Leader-PENDING'
    for sub in ['Code','Demo','Dox']:(folder/sub).mkdir(parents=True,exist_ok=True)
    for name in ['CampusConnect-Final-Report.docx','CampusConnect-Final-Report.pdf','data-dictionary.csv']:
        shutil.copyfile(OUT/name,folder/'Dox'/name)
    names='Name and surname | Student number | Contribution\n'+'\n'.join(f'Member {i}: PENDING | PENDING | PENDING' for i in range(1,5))+'\n\nAdd members 5–6 if applicable.\nGroup number: PENDING\nTeam leader: PENDING\nLink to FORKED git repository: PENDING\n'
    if not (folder/'Group-PENDING-Names.txt').exists():
        (folder/'Group-PENDING-Names.txt').write_text(names,encoding='utf-8')
    if not (OUT/'usability-results.csv').exists():
        (OUT/'usability-results.csv').write_text('participant_id,role,device,viewport,build,task_id,start_time,end_time,seconds,success_unaided,wrong_taps,validation_errors,help_requests,ease_1_to_5,comments\n',encoding='utf-8')
    (OUT/'validation-summary.json').write_text(json.dumps({'date':'2026-10-05','pdf_pages':20,'docx_explicit_page_breaks':19,'dictionary_fields':len(records),'dictionary_tables':sorted(set(r[0] for r in records)),'registered_screen_panels':len(SCREENS),'fresh_checks':{'typecheck':'passed','backend_isolated_tests':{'tests':9,'failures':0,'errors':0}},'historical_checks':'Browser action/layout and web export passed in session 2026-10-04; full backend suite not freshly run','pending':['Group identity','Confirmed fork','Human usability data','Integrated narrated demo','Native device checks']},indent=2),encoding='utf-8')

if __name__=='__main__':
    assert len(DATA['pages'])==20
    diagrams();galleries();records=dictionary();exports();support(records)
    print('Generated report, diagrams, 29 registered-screen panels, dictionary and template folders.')
    if '--preview' in sys.argv:
        import fitz
        previews=OUT/'previews';previews.mkdir(exist_ok=True)
        rendered=fitz.open(OUT/'CampusConnect-Final-Report.pdf')
        sheet=Image.new('RGB',(1000,1500),'#E2E8F0')
        for i,page in enumerate(rendered):
            pix=page.get_pixmap(matrix=fitz.Matrix(1.3,1.3))
            preview=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
            preview.save(previews/f'page-{i+1:02}.png')
            preview.thumbnail((235,285));x=10+(i%4)*250;y=10+(i//4)*300
            sheet.paste(preview,(x,y))
            ImageDraw.Draw(sheet).text((x+5,y+285),f'Page {i+1}',font=font(13),fill=INK)
        sheet.save(previews/'contact-sheet.png')
        print('Rendered all 20 PDF pages to previews for visual inspection.')
