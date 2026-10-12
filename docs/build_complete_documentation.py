"""Build consolidated, source-based project documentation without accessing live records."""
from pathlib import Path
import re, sys, json, zipfile, hashlib
from html import escape

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.report-tools'))
from docx import Document
from docx.shared import Pt, Inches
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from pypdf import PdfReader

OUT = ROOT / 'docs/complete-documentation'
OUT.mkdir(parents=True, exist_ok=True)
JAVA = ROOT / 'backend/src/main/java/com/campusconnect'
parts = []
def read(p): return p.read_text(encoding='utf-8-sig', errors='replace')
def scrub(s):
    s = re.sub(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', '<redacted-email>', s)
    s = re.sub(r'\$2[aby]\$\d{2}\$[A-Za-z0-9./]+', '<redacted-password-hash>', s)
    s = re.sub(r'eyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+', '<redacted-token>', s)
    s = re.sub(r'(?im)^(\s*(?:password|secret|username|from)\s*:\s*).+$', r'\1<configure externally>', s)
    return s
def add(title, body): parts.append('# ' + title + '\n\n' + scrub(body).strip() + '\n')
def table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---']*len(headers)) + ' |'] + ['| ' + ' | '.join(str(x).replace('|','/').replace('\n',' ') for x in row) + ' |' for row in rows])

add('CampusConnect — Complete Project Documentation', '''Prepared: 11 October 2026. Project: Campus-Ride-Connect / CampusConnect.

This document consolidates the current local implementation, existing technical guides, screen catalogue, API declarations, Java data mappings, MySQL schema and migration history, test evidence, and maintenance material. It is intended for project handover, study, review, and further development.

Evidence boundary: code presence describes implementation, not successful execution of every feature. Earlier testing results are explicitly historical. No application tests, live database inspection, database export, payment, email, or emergency dispatch were performed during this documentation build. The included database package contains source definitions, not a backup of an installed MySQL database. Credentials, account exports, tokens, identity files and private audio are excluded.

Read sections 1–12 for the project guide; use the remaining sections as detailed references. Source-relative paths refer to the inner Campus-Ride-Connect project directory.''')
add('1. Project purpose, scope and current status', '''CampusConnect coordinates university transport and safety for riders, drivers, administrators and campus security. The interface is configured around Sol Plaatje University in Kimberley, South Africa, and uses rand-denominated fares.

Implemented application paths cover accounts and password recovery, driver onboarding and approval, immediate and scheduled ride requests, trip status and location updates, ride messages, trip history and ratings, payment display metadata and pending payment records, SOS alerts with evidence, administration and security review.

The repository is a full-stack prototype. Public hosting, physical-device validation, payment settlement, biometric matching, production security hardening, and a consolidated clean database installation remain incomplete or unverified. No participant study or operational performance measurements are supplied.''')
add('2. Roles and practical user workflows', table(['Role','Workflow'], [
('Rider','Register/sign in; choose pickup and destination; review estimate; request or schedule; wait for driver; follow trip; chat; report SOS if needed; complete and rate; inspect history and account/payment metadata.'),
('Driver','Submit account and vehicle/application information; await approval; switch availability; inspect pending requests; accept/decline; progress enroute, arrived, started and completed stages; update location; chat; use SOS; inspect history and earnings.'),
('Administrator','Sign in with privileged account; inspect statistics; review driver applications; approve/reject; manage user status; monitor rides; inspect incidents/audio; update university settings.'),
('Security','Sign in with privileged account; monitor active rides; inspect SOS context and evidence; record dispatch action; inspect resolved incidents.')]) + '''

The ride enum defines SCHEDULED, PENDING, ACCEPTED, ENROUTE, ARRIVED, STARTED, COMPLETED and CANCELLED. This is a logical lifecycle, not proof that all transitions are strictly enforced. RideService derives scheduled status from the requested status and scheduledAt, copies supplied journey/fare fields, persists the ride, and records start/completion timestamps when status changes. Fare integrity and concurrent acceptance require integration review.

Payment selection stores metadata; a payment API creates a PENDING record and does not charge funds. Face capture/upload does not perform biometric matching. SOS dispatch is a recorded application action; actual emergency response has not been validated.''')
for number, name in [(3,'architecture'),(4,'requirements'),(5,'deployment'),(6,'api-spec'),(7,'security'),(8,'testing'),(9,'responsive-layout'),(10,'password-reset-email')]:
    p = ROOT / f'docs/{name}.md'
    add(f'{number}. {name.replace("-", " ").title()}', f'Source: `{p.relative_to(ROOT).as_posix()}`\n\n' + read(p))
add('11. Screen catalogue and demonstration', read(ROOT/'docs/report/screen-catalogue.md') + '\n\n' + read(ROOT/'docs/report/demo-script.md'))
add('12. Database model and installation boundaries', '''MySQL is the application store; Spring Data JPA supplies entity/repository persistence and JdbcTemplate performs selected direct SQL operations. Development configuration uses Hibernate ddl-auto:update, which can create or modify mapped tables. Some controller startup logic also creates or repairs structures. These mechanisms are not a versioned migration strategy.

Key logical relationships: users have driver profiles; rides belong to riders and may be assigned to driver users; messages belong to rides and senders; SOS records belong to rides and reporting users, with optional resolving users; saved payment methods belong to users; payment records refer to rides and payers. Some Java references are scalar IDs, so referential constraints must be verified in installed DDL.

Important drift: schema.sql declares rides.id and duration_minutes while Ride.java declares ride_id and duration_min. The entity also has scheduled/start/completion timestamps and current coordinates not fully represented by the base script. The schema driver reference targets users, while an older migration requires a correction of driver references. Historical migration prefixes overlap, including two 001 scripts. Some scripts expect a preexisting schema and are not safe to rerun. fix_rides_table.sql rebuilds rides. The Payment Java source is a stub while PaymentController inserts into payments; complete payments creation DDL is not evidenced by the existing dictionary. Empty files are listed explicitly later.

Use the dictionary and entity declarations together: SQL types reflect authored scripts; Java types and annotations reflect ORM mappings; neither establishes the installed database structure. Before installation or upgrades, inspect SHOW CREATE TABLE, choose one reconciled schema/migration baseline, use disposable data, and test migration and restore behavior. Do not run the appended SQL files as one installation script.

For a live schema reference, an operator can use SHOW TABLES; SHOW CREATE TABLE table_name; SHOW INDEX FROM table_name; and query INFORMATION_SCHEMA.KEY_COLUMN_USAGE. For a future schema-only backup use mysqldump --no-data --routines --triggers --events --user=<database-user> --password --result-file=campus_connect_schema.sql campus_connect. The password flag prompts rather than embedding a secret. This command was not executed here. Row backups require separate private storage and a verified restore procedure.''')
add('13. Detailed data dictionary', read(ROOT/'docs/report/data-dictionary.md'))
add('13A. Design diagrams and captured interfaces', '''The following are existing project report diagrams and supplied interface captures. They illustrate the documented design and captured states, rather than a newly verified installed database or current running session.

![Database relationship design](../report/assets/erd.png)

![Class design](../report/assets/class.png)

![Ride lifecycle design](../report/assets/state.png)

![Ride request sequence](../report/assets/sequence.png)

![Rider booking capture](../images/rider-booking.png)

![Administrator dashboard capture](../images/admin-dashboard.png)

![Driver earnings capture](../images/driver-earnings.png)''')

entity_count = 0
for p in sorted((JAVA/'entity').glob('*.java')):
    s = read(p); entity_count += 1
    fields = re.findall(r'((?:\s*@[^\n]+\n)*)\s*private\s+([\w<>., ?]+)\s+(\w+)\s*(?:=[^;]*)?;',s)
    rows=[]
    for annotations, typ, name in fields:
        rows.append((name, typ.strip(), ' '.join(annotations.split()) or 'No explicit field annotation'))
    body=f'Source: `{p.relative_to(ROOT).as_posix()}`. Java types and annotations are logical mappings; actual MySQL DDL has not been inspected.\n\n'
    body += table(['Field','Java type','Mapping / constraints'],rows) if rows else 'No declared persistent fields found. This file may be an unimplemented placeholder.\n\n```java\n' + s.strip() + '\n```'
    enums = re.findall(r'(?:public\s+)?enum\s+(\w+)\s*\{([^}]+)\}',s,re.S)
    for name, values in enums: body+='\n\nEnum '+name+': '+ ' '.join(values.split())
    add('14.'+str(entity_count)+' Entity: '+p.stem,body)

routes=[]
for p in sorted((JAVA/'controller').glob('*.java')):
    s=read(p)
    before=s.split('public class')[0]
    base=re.search(r'@RequestMapping\s*\(\s*"([^"]+)"',before)
    prefix=base.group(1) if base else ''
    for m in re.finditer(r'@(Get|Post|Put|Delete|Patch|Request)Mapping\b(?:\(([^\n]*)\))?',s):
        if m.start()<len(before): continue
        args=m.group(2) or ''; paths=re.findall(r'"([^"]*)"',args)
        path=paths[0] if paths and ('/' in paths[0] or not paths[0]) else ''
        tail=s[m.end():]
        signature=re.search(r'public\s+[^;{]+\(',tail)
        methodname=re.search(r'(\w+)\s*\($',signature.group(0)).group(1) if signature else 'inspect source'
        routes.append((m.group(1).upper().replace('REQUEST','REQUEST*'),prefix+path,p.stem,methodname))
add('15. Complete controller route declaration index', 'Extracted from current Java annotations. RequestMapping entries marked REQUEST* may require additional method/path inspection. This index does not infer role enforcement, request schemas, alternate mapping arrays or successful live responses. The detailed controller contract appendix below supplies the declaration context.\n\n'+table(['HTTP','Declared path','Controller','Handler'],routes))

for folder in ['controller','dto','repository','service','config','security','validation','exception','mapper','util','websocket']:
    rows=[]
    for p in sorted((JAVA/folder).glob('*.java')):
        s=read(p)
        methods=re.findall(r'(?:public|protected)\s+(?:static\s+)?[\w<>?, .\[\]]+\s+(\w+)\s*\(',s)
        imports=re.findall(r'import com\.campusconnect\.([^;]+);',s)
        rows.append((p.stem,', '.join(dict.fromkeys(methods)) or 'No public/protected method declarations',', '.join(imports) or 'No project imports'))
    add('16. Backend module inventory: '+folder,table(['Module','Declared operations','Project dependencies'],rows))

for folder in ['controller','dto','repository']:
    for p in sorted((JAVA/folder).glob('*.java')):
        s=read(p)
        if folder=='controller':
            # Declaration context only: avoid copying startup credentials and account seed bodies.
            declarations=[]
            for m in re.finditer(r'@(?:Get|Post|Put|Delete|Patch|Request)Mapping\b',s):
                end=s.find('{',m.start())
                if end>=0: declarations.append(s[m.start():end].strip())
            s='\n\n'.join(declarations)
        elif folder=='dto':
            s='\n'.join(line for line in s.splitlines() if line.strip().startswith(('@','private ','public enum')))
        add('17. Contract reference: '+p.stem, f'Source: `{p.relative_to(ROOT).as_posix()}`\n\n```java\n{s.strip() or "Empty or placeholder source file."}\n```')

rows=[]
for folder in ['screens','components','services','context','hooks','navigation','utils','theme']:
    for p in sorted((ROOT/'mobile/src'/folder).glob('*')):
        if not p.is_file(): continue
        s=read(p)
        deps=re.findall(r'import[\s\S]*?from\s+[\'"]([^\'"]+)[\'"]',s)
        rows.append((p.relative_to(ROOT).as_posix(),len(s.splitlines()),', '.join(dict.fromkeys(deps))))
add('18. Complete client module inventory', 'Screens, shared UI, network/storage services, state providers, hooks, navigation, validation and theme modules are indexed below. Route registration details are in the screen catalogue. File presence alone does not establish that a screen is reachable.\n\n'+table(['Source file','Lines','Imports / dependencies'],rows))
for folder in ['services','context','hooks']:
    rows=[]
    for p in sorted((ROOT/'mobile/src'/folder).glob('*')):
        if not p.is_file(): continue
        s=read(p)
        operations=re.findall(r'(?:async\s+)?(\w+)\s*:\s*async|(?:export\s+)?(?:async\s+)?function\s+(\w+)|(?:public\s+)?(?:static\s+)?async\s+(\w+)\s*\(',s)
        endpoints=re.findall(r'(?:apiClient|api)\.(get|post|put|delete|patch)\s*(?:<[^>]+>)?\s*\(\s*([`\'"][^\n]+)',s)
        rows.append((p.stem,', '.join(next(x for x in tup if x) for tup in operations),'; '.join(verb.upper()+' '+expr for verb,expr in endpoints)))
    add('19. Client operations: '+folder,table(['Module','Named async/functions','API call source expressions'],rows))
pkg=json.loads(read(ROOT/'mobile/package.json'))
add('20. Dependency and command reference', 'Client manifest: `mobile/package.json`. Version ranges below are authored manifest declarations, not a fresh resolved installation.\n\n'+table(['Command','Implementation'],pkg['scripts'].items())+'\n\n'+table(['Client dependency','Version declaration'],list(pkg['dependencies'].items())+list(pkg['devDependencies'].items()))+'\n\nBackend manifest:\n\n```xml\n'+read(ROOT/'backend/pom.xml')+'\n```\n\nRoot aggregator manifest:\n\n```xml\n'+read(ROOT/'pom.xml')+'\n```')
for p in sorted((ROOT/'backend/src/main/resources').glob('application*.yml')):
    add('21. Sanitized configuration: '+p.name,'Source: `'+p.relative_to(ROOT).as_posix()+'`\n\n```yaml\n'+read(p)+'\n```')
for p in sorted((ROOT/'backend/docker').glob('*')):
    if p.is_file(): add('21. Container configuration: '+p.name,'```\n'+read(p)+'\n```')
tests=[]
for p in sorted((ROOT/'backend/src/test').rglob('*.java')):
    s=read(p); names=re.findall(r'@Test\s+(?:@[\s\S]*?\n\s*)?(?:public\s+)?void\s+(\w+)\s*\(',s)
    tests.append((p.relative_to(ROOT).as_posix(),', '.join(names) or 'Inspect source for test declarations'))
for p in sorted((ROOT/'mobile/scripts').glob('*.cjs')): tests.append((p.relative_to(ROOT).as_posix(),'Browser/fixture check; inspect script prerequisites'))
for p in sorted(ROOT.glob('test*.ps1')): tests.append((p.name,'PowerShell API workflow script; inspect credentials/data writes before running'))
add('22. Test source inventory and evaluation plan',table(['Test source','Cases / purpose'],tests)+'\n\n'+read(ROOT/'docs/report/usability-study.md')+'\n\n'+read(ROOT/'docs/report/submission-checklist.md'))
history=['ADMIN_DASHBOARD_FIX.md','DRIVER_FIX_CHECKLIST.md','DRIVER_REQUEST_API_FIX.md','DRIVER_SCREEN_ENHANCEMENTS.md','FINDING_DRIVERS_SCREEN_IMPLEMENTATION.md','RIDE_REQUEST_IMPLEMENTATION.md','IMPLEMENTATION_SUMMARY.md','BACKEND_API_REQUIREMENTS.md']
add('23. Implementation history index', 'These earlier development notes record work, requirements and proposed fixes. Their claims must be checked against current source; they do not supersede the current status and limitations in this document.\n\n'+table(['Note','Opening description'],[(name,' '.join(read(ROOT/name).splitlines()[0:4])) for name in history if (ROOT/name).exists()]))

dbrows=[]
DBOUT=OUT/'database-reference'; DBOUT.mkdir(exist_ok=True)
for p in sorted((ROOT/'database/mysql').rglob('*.sql')):
    s=read(p); rel=p.relative_to(ROOT/'database/mysql')
    # Data-only seed/maintenance scripts are described, not distributed with identifying rows.
    dataonly = p.name in ['seed.sql','004_add_pending_driver_for_approval.sql'] or 'maintenance' in p.parts
    ddl = re.sub(r'(?is)\bINSERT\s+INTO\b.*?;', '-- Data insertion omitted from documentation; use controlled private fixtures.',s)
    exported = '-- Data-only source omitted: account identifiers, names, credentials and row values are private.\n' if dataonly else scrub(ddl)
    dest=DBOUT/rel; dest.parent.mkdir(parents=True,exist_ok=True); dest.write_text(exported,encoding='utf-8')
    statements=re.findall(r'(?im)^\s*(CREATE\s+(?:TABLE|VIEW|PROCEDURE)|ALTER\s+TABLE|DROP\s+TABLE|UPDATE|INSERT\s+INTO|DELETE\s+FROM)\s+(?:IF\s+(?:NOT\s+)?EXISTS\s+)?([^\s(]+)',s)
    dbrows.append((rel.as_posix(),len(s.splitlines()),'; '.join(a+' '+b for a,b in statements) or 'Empty/comment-only or no matched statement', 'Data-only omitted' if dataonly else 'Sanitized structure/reference'))
    add('24. Database source: '+rel.as_posix(), f'Source: `database/mysql/{rel.as_posix()}`. Historical script, not an approved execution order.\n\n```sql\n{exported.strip() or "-- Empty source file; no definitions implemented."}\n```')
add('25. Database migration and maintenance catalogue',table(['Script','Lines','Declared operations','Package treatment'],dbrows))
(DBOUT/'README.md').write_text('''# Source database reference

This folder contains sanitized copies of the repository SQL definitions and historical migrations, not a live database dump or a reconciled installer. Account insertions and data-only seed/maintenance bodies are omitted. Some source files are empty. Do not execute scripts in filename order: migration prefixes overlap, schemas drift, and fix_rides_table.sql rebuilds a table. Refer to the complete documentation for entity mappings and installation caveats. No database was read or changed during generation.
''',encoding='utf-8')

inventory=[]
skip={'node_modules','.git','.report-tools','target','dist','.expo','submission','complete-documentation','__pycache__'}
for p in sorted(ROOT.rglob('*')):
    if not p.is_file() or any(x in skip for x in p.relative_to(ROOT).parts): continue
    rel=p.relative_to(ROOT).as_posix()
    inventory.append((rel,p.stat().st_size))
add('26. Project file inventory', 'Generated files, vendored dependencies, build outputs, Git metadata and duplicate submission packages are omitted. Sensitive filenames are listed only; their contents are excluded. Screenshots and earlier report assets are existing evidence, not newly captured runtime states.\n\n'+table(['Project-relative path','Bytes'],inventory))
add('27. Handover, unresolved work and document provenance', '''To continue the project, first reconcile SQL and Java mappings and validate a clean installation. Restrict privileged registration and startup account routines, redact account responses, review ownership and concurrent acceptance, and externalize secrets. Then verify full trips, email recovery, native camera/location/audio permissions, upload access and retention, and device sharing. Payment and biometric providers need explicit integration design. Establish HTTPS hosting, backups and restore verification, controlled logs and deployment configuration before release.

Existing academic report assets and submission files remain intact. Group identity, contribution records, repository confirmation, actual participant results and presentation evidence remain the project owner's supplied information; this document does not fill those fields with invented details.

This report was built from local source and authored documentation on 11 October 2026. Earlier statements of 4–5 October testing are retained as historical evidence, not rerun results. No external research was needed for the source inventory. Repeat generation with python docs/build_complete_documentation.py. The Markdown is the searchable master; Word and PDF are reading formats; database-reference is the sanitized SQL reference companion.''')

md='\n\n'.join(parts)
# Provide an explicit contents list using the actual generated section headings.
titles=[x.splitlines()[0][2:] for x in parts]
contents='# Contents\n\n'+'\n'.join('- '+x for x in titles[1:])+'\n\n'
md=parts[0]+'\n'+contents+'\n\n'.join(parts[1:])
stem='CampusConnect-Complete-Documentation'
(OUT/(stem+'.md')).write_text(md,encoding='utf-8')
doc=Document(); doc.sections[0].header.paragraphs[0].text='CampusConnect | Complete Project Documentation | 11 October 2026'
doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10)
styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name='SourceCode',fontName='Courier',fontSize=6.5,leading=8,wordWrap='CJK'))
story=[]; code=False; lines=md.splitlines(); i=0
while i<len(lines):
    line=lines[i]; i+=1
    if line.startswith('```'): code=not code; continue
    if not line.strip(): story.append(Spacer(1,4)); continue
    if line.startswith('#') and not code:
        level=min(len(line)-len(line.lstrip('#')),3); value=line.lstrip('# ').strip()
        doc.add_heading(value,level=level)
        if level==1 and story: story.append(PageBreak())
        story.append(Paragraph(escape(value),styles['Heading'+str(level)])); continue
    if re.match(r'^\|[\s:|\-]+\|$',line): continue
    figure=re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)',line.strip())
    if figure and not code:
        asset=(OUT/figure.group(2)).resolve()
        if asset.is_file():
            from PIL import Image as PILImage
            with PILImage.open(asset) as im: width,height=im.size
            ratio=min(480/width,620/height)
            doc.add_picture(str(asset),width=Inches(min(6.3,width*ratio/72)))
            doc.add_paragraph(figure.group(1))
            story.append(Image(str(asset),width=width*ratio,height=height*ratio))
            story.append(Paragraph(escape(figure.group(1)),styles['BodyText']))
            continue
    value=line
    if not code: value=re.sub(r'!\[([^]]*)\]\(([^)]+)\)',r'[Figure: \1; source asset: \2]',value); value=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1 (\2)',value); value=value.replace('**','').replace('`','')
    p=doc.add_paragraph(value)
    if code:
        for r in p.runs: r.font.name='Consolas'; r.font.size=Pt(7)
    # Paragraph cells wrap consistently, avoiding page-wide tables in source appendices.
    story.append(Paragraph(escape(value),styles['SourceCode' if code else 'BodyText']))
    story.append(Spacer(1,3))
doc.save(OUT/(stem+'.docx'))
def footer(canvas, document):
    canvas.setFont('Helvetica',8); canvas.drawString(36,22,'CampusConnect | Source-based documentation | 11 October 2026'); canvas.drawRightString(A4[0]-36,22,str(document.page))
SimpleDocTemplate(str(OUT/(stem+'.pdf')),pagesize=A4,rightMargin=36,leftMargin=36,topMargin=38,bottomMargin=36,title='CampusConnect Complete Project Documentation',author='CampusConnect Project').build(story,onFirstPage=footer,onLaterPages=footer)
with zipfile.ZipFile(OUT/'CampusConnect-Complete-Documentation-Package.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p in sorted(OUT.rglob('*')):
        if p.is_file() and p.suffix!='.zip': z.write(p,p.relative_to(OUT))
    for p in [ROOT/'docs/report/assets'/name for name in ['erd.png','class.png','state.png','sequence.png']]+[ROOT/'docs/images'/name for name in ['rider-booking.png','admin-dashboard.png','driver-earnings.png']]:
        z.write(p,'reference-assets/'+p.name)
result={'pages':len(PdfReader(OUT/(stem+'.pdf')).pages),'routes':len(routes),'entities':entity_count,'sql_files':len(dbrows),'inventory_files':len(inventory),'words':len(md.split()),'outputs':[p.name for p in OUT.glob('*') if p.is_file()]}
(OUT/'generation-summary.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
